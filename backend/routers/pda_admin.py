from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Request
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from typing import Dict, List, Optional
import os
import io
import csv
from io import BytesIO
from pathlib import Path
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

from database import get_db
from models import (
    PdaItem,
    PdaTeam,
    PdaGallery,
    PdaUser,
    PdaAdmin,
    PdaEventRegistration,
    PdaEventTeam,
    PdaEventTeamMember,
    PdaEventAttendance,
    PdaEventScore,
    PdaEventBadge,
    PdaEventInvite,
    PersohubCommunity,
    PersohubCommunityFollow,
    PersohubPost,
    PersohubPostAttachment,
    PersohubPostLike,
    PersohubPostComment,
    PersohubPostHashtag,
    PersohubPostMention,
    PersohubHashtag,
)
from schemas import (
    ProgramCreate, ProgramUpdate, ProgramResponse,
    EventCreate, EventUpdate, EventResponse,
    PdaTeamCreate, PdaTeamUpdate, PdaTeamResponse,
    PdaAdminUserUpdate, PdaAdminUserResponse,
    PdaGalleryCreate, PdaGalleryUpdate, PdaGalleryResponse,
    PresignRequest, PresignResponse,
    PdaPdfPreviewGenerateRequest, PdaPdfPreviewGenerateResponse,
)
from security import require_pda_home_admin, require_superadmin
from utils import (
    S3_BUCKET_NAME,
    S3_CLIENT,
    _build_s3_url,
    _extract_s3_key_from_url,
    _generate_presigned_put_url,
    _upload_to_s3,
    log_admin_action,
)
from auth import get_password_hash
from recruitment_state import get_recruitment_state, get_recruitment_state_map
from persohub_service import is_profile_name_valid, generate_unique_profile_name

try:
    import pypdfium2 as pdfium
except Exception:  # pragma: no cover
    pdfium = None

router = APIRouter()

PDF_PREVIEW_MAX_PAGES = 20
PDF_PREVIEW_RENDER_SCALE = 1.5


def _render_pdf_preview_images(pdf_bytes: bytes, max_pages: int) -> List[bytes]:
    if not pdfium:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="PDF preview dependency missing (install pypdfium2)",
        )
    try:
        doc = pdfium.PdfDocument(pdf_bytes)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid PDF file") from exc

    page_count = len(doc)
    target_pages = min(max_pages, PDF_PREVIEW_MAX_PAGES, page_count)
    images: List[bytes] = []

    for idx in range(target_pages):
        page = doc[idx]
        bitmap = None
        pil_image = None
        output = BytesIO()
        try:
            bitmap = page.render(scale=PDF_PREVIEW_RENDER_SCALE)
            pil_image = bitmap.to_pil()
            pil_image.save(output, format="WEBP", quality=84, method=6)
            images.append(output.getvalue())
        finally:
            output.close()
            if pil_image is not None:
                try:
                    pil_image.close()
                except Exception:
                    pass
            if bitmap is not None and hasattr(bitmap, "close"):
                try:
                    bitmap.close()
                except Exception:
                    pass
            if hasattr(page, "close"):
                try:
                    page.close()
                except Exception:
                    pass

    if hasattr(doc, "close"):
        doc.close()
    return images

def _build_team_response(member: PdaTeam, user: Optional[PdaUser], resume_url: Optional[str] = None) -> PdaTeamResponse:
    return PdaTeamResponse(
        id=member.id,
        user_id=member.user_id,
        name=user.name if user else None,
        profile_name=user.profile_name if user else None,
        regno=user.regno if user else None,
        dept=user.dept if user else None,
        email=user.email if user else None,
        phno=user.phno if user else None,
        dob=user.dob if user else None,
        resume_url=resume_url,
        team=member.team,
        designation=member.designation,
        photo_url=user.image_url if user else None,
        instagram_url=user.instagram_url if user else None,
        linkedin_url=user.linkedin_url if user else None,
        github_url=user.github_url if user else None,
        created_at=member.created_at
    )


def _build_admin_user_response(
    user: PdaUser,
    member: Optional[PdaTeam],
    recruit: Optional[Dict[str, Optional[str]]] = None,
) -> PdaAdminUserResponse:
    recruit_state = recruit or {}
    return PdaAdminUserResponse(
        id=user.id,
        team_member_id=member.id if member else None,
        name=user.name,
        profile_name=user.profile_name,
        regno=user.regno,
        dept=user.dept,
        email=user.email,
        phno=user.phno,
        dob=user.dob,
        gender=user.gender,
        is_member=bool(user.is_member),
        is_applied=bool(recruit_state.get("is_applied")),
        preferred_team=recruit_state.get("preferred_team"),
        preferred_team_1=recruit_state.get("preferred_team_1"),
        preferred_team_2=recruit_state.get("preferred_team_2"),
        preferred_team_3=recruit_state.get("preferred_team_3"),
        resume_url=recruit_state.get("resume_url"),
        email_verified=bool(user.email_verified_at),
        team=member.team if member else None,
        designation=member.designation if member else None,
        photo_url=user.image_url,
        instagram_url=user.instagram_url,
        linkedin_url=user.linkedin_url,
        github_url=user.github_url,
        created_at=user.created_at,
    )


def _sync_member_status_from_team(user: Optional[PdaUser], member: Optional[PdaTeam]) -> None:
    if not user or not member:
        return
    # Team assignment is considered complete only when both values exist.
    if member.team and member.designation:
        user.is_member = True


def _user_dependency_checks(db: Session, user_id: int) -> List[str]:
    checks = [
        ("pda_event_registrations", db.query(PdaEventRegistration).filter(PdaEventRegistration.user_id == user_id).first()),
        ("pda_event_teams", db.query(PdaEventTeam).filter(PdaEventTeam.team_lead_user_id == user_id).first()),
        ("pda_event_team_members", db.query(PdaEventTeamMember).filter(PdaEventTeamMember.user_id == user_id).first()),
        ("pda_event_attendance", db.query(PdaEventAttendance).filter(
            (PdaEventAttendance.user_id == user_id) | (PdaEventAttendance.marked_by_user_id == user_id)
        ).first()),
        ("pda_event_scores", db.query(PdaEventScore).filter(PdaEventScore.user_id == user_id).first()),
        ("pda_event_badges", db.query(PdaEventBadge).filter(PdaEventBadge.user_id == user_id).first()),
        ("pda_event_invites", db.query(PdaEventInvite).filter(
            (PdaEventInvite.invited_user_id == user_id) | (PdaEventInvite.invited_by_user_id == user_id)
        ).first()),
        ("persohub_communities", db.query(PersohubCommunity).filter(PersohubCommunity.admin_id == user_id).first()),
        ("persohub_community_follows", db.query(PersohubCommunityFollow).filter(PersohubCommunityFollow.user_id == user_id).first()),
        ("persohub_posts", db.query(PersohubPost).filter(PersohubPost.admin_id == user_id).first()),
        ("persohub_post_likes", db.query(PersohubPostLike).filter(PersohubPostLike.user_id == user_id).first()),
        ("persohub_post_comments", db.query(PersohubPostComment).filter(PersohubPostComment.user_id == user_id).first()),
        ("persohub_post_mentions", db.query(PersohubPostMention).filter(PersohubPostMention.user_id == user_id).first()),
    ]
    return [name for name, hit in checks if hit is not None]


@router.post("/pda-admin/programs", response_model=ProgramResponse)
def create_pda_program(
    program_data: ProgramCreate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    new_program = PdaItem(
        type="program",
        title=program_data.title,
        description=program_data.description,
        tag=program_data.tag,
        poster_url=program_data.poster_url,
        start_date=program_data.start_date,
        end_date=program_data.end_date,
        format=program_data.format,
        hero_url=program_data.hero_url,
        featured_poster_url=program_data.featured_poster_url,
        is_featured=program_data.is_featured
    )
    db.add(new_program)
    db.commit()
    db.refresh(new_program)
    log_admin_action(db, admin, "Create PDA program", request.method if request else None, request.url.path if request else None, {"program_id": new_program.id})
    return ProgramResponse.model_validate(new_program)


@router.put("/pda-admin/programs/{program_id}", response_model=ProgramResponse)
def update_pda_program(
    program_id: int,
    program_data: ProgramUpdate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    program = db.query(PdaItem).filter(PdaItem.id == program_id, PdaItem.type == "program").first()
    if not program:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Program not found")

    if program_data.title is not None:
        program.title = program_data.title
    if program_data.description is not None:
        program.description = program_data.description
    if program_data.tag is not None:
        program.tag = program_data.tag
    if program_data.poster_url is not None:
        program.poster_url = program_data.poster_url
    if program_data.start_date is not None:
        program.start_date = program_data.start_date
    if program_data.end_date is not None:
        program.end_date = program_data.end_date
    if program_data.format is not None:
        program.format = program_data.format
    if program_data.hero_url is not None:
        program.hero_url = program_data.hero_url
    if program_data.featured_poster_url is not None:
        program.featured_poster_url = program_data.featured_poster_url
    if program_data.is_featured is not None:
        program.is_featured = program_data.is_featured

    db.commit()
    db.refresh(program)
    log_admin_action(db, admin, "Update PDA program", request.method if request else None, request.url.path if request else None, {"program_id": program_id})
    return ProgramResponse.model_validate(program)


@router.delete("/pda-admin/programs/{program_id}")
def delete_pda_program(
    program_id: int,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    program = db.query(PdaItem).filter(PdaItem.id == program_id, PdaItem.type == "program").first()
    if not program:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Program not found")
    db.delete(program)
    db.commit()
    log_admin_action(db, admin, "Delete PDA program", request.method if request else None, request.url.path if request else None, {"program_id": program_id})
    return {"message": "Program deleted successfully"}


@router.post("/pda-admin/home-events", response_model=EventResponse)
def create_pda_event(
    event_data: EventCreate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    new_event = PdaItem(
        type="event",
        title=event_data.title,
        start_date=event_data.start_date,
        end_date=event_data.end_date,
        format=event_data.format,
        description=event_data.description,
        poster_url=event_data.poster_url,
        hero_url=event_data.hero_url,
        featured_poster_url=event_data.featured_poster_url,
        is_featured=event_data.is_featured
    )
    db.add(new_event)
    db.commit()
    db.refresh(new_event)
    log_admin_action(db, admin, "Create PDA event", request.method if request else None, request.url.path if request else None, {"event_id": new_event.id})
    return EventResponse.model_validate(new_event)


@router.put("/pda-admin/home-events/{event_id}", response_model=EventResponse)
def update_pda_event(
    event_id: int,
    event_data: EventUpdate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    event = db.query(PdaItem).filter(PdaItem.id == event_id, PdaItem.type == "event").first()
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found")

    if event_data.title is not None:
        event.title = event_data.title
    if event_data.start_date is not None:
        event.start_date = event_data.start_date
    if event_data.end_date is not None:
        event.end_date = event_data.end_date
    if event_data.format is not None:
        event.format = event_data.format
    if event_data.description is not None:
        event.description = event_data.description
    if event_data.poster_url is not None:
        event.poster_url = event_data.poster_url
    if event_data.hero_url is not None:
        event.hero_url = event_data.hero_url
    if event_data.featured_poster_url is not None:
        event.featured_poster_url = event_data.featured_poster_url
    if event_data.is_featured is not None:
        event.is_featured = event_data.is_featured

    db.commit()
    db.refresh(event)
    log_admin_action(db, admin, "Update PDA event", request.method if request else None, request.url.path if request else None, {"event_id": event_id})
    return EventResponse.model_validate(event)


@router.post("/pda-admin/home-events/{event_id}/feature", response_model=EventResponse)
def feature_pda_event(
    event_id: int,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    event = db.query(PdaItem).filter(PdaItem.id == event_id, PdaItem.type == "event").first()
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found")

    event.is_featured = True
    db.commit()
    db.refresh(event)
    log_admin_action(db, admin, "Feature PDA event", request.method if request else None, request.url.path if request else None, {"event_id": event_id})
    return EventResponse.model_validate(event)


@router.delete("/pda-admin/home-events/{event_id}")
def delete_pda_event(
    event_id: int,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    event = db.query(PdaItem).filter(PdaItem.id == event_id, PdaItem.type == "event").first()
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found")
    db.delete(event)
    db.commit()
    log_admin_action(db, admin, "Delete PDA event", request.method if request else None, request.url.path if request else None, {"event_id": event_id})
    return {"message": "Event deleted successfully"}


@router.get("/pda-admin/team", response_model=List[PdaTeamResponse])
def list_team_members(
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db)
):
    rows = (
        db.query(PdaTeam, PdaUser)
        .join(PdaUser, PdaTeam.user_id == PdaUser.id, isouter=True)
        .order_by(PdaTeam.team.asc().nullslast(), PdaTeam.designation.asc().nullslast(), PdaUser.name.asc().nullslast())
        .all()
    )
    users = [u for _, u in rows if u and u.id is not None]
    resume_map = get_recruitment_state_map(db, users) if users else {}
    return [
        _build_team_response(
            m,
            u,
            (resume_map.get(u.id) or {}).get("resume_url") if u else None,
        )
        for m, u in rows
    ]


@router.get("/pda-admin/users", response_model=List[PdaAdminUserResponse])
def list_pda_users(
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
):
    users = (
        db.query(PdaUser)
        .order_by(PdaUser.name.asc().nullslast(), PdaUser.regno.asc())
        .all()
    )
    user_ids = [user.id for user in users]
    team_rows = (
        db.query(PdaTeam)
        .filter(PdaTeam.user_id.in_(user_ids))
        .order_by(PdaTeam.id.asc())
        .all()
        if user_ids else []
    )
    team_map = {}
    for member in team_rows:
        if member.user_id not in team_map:
            team_map[member.user_id] = member
    recruit_map = get_recruitment_state_map(db, users)
    return [_build_admin_user_response(user, team_map.get(user.id), recruit_map.get(user.id)) for user in users]


@router.put("/pda-admin/users/{user_id}", response_model=PdaAdminUserResponse)
def update_pda_user_admin(
    user_id: int,
    user_data: PdaAdminUserUpdate,
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    user = db.query(PdaUser).filter(PdaUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    updates = user_data.model_dump(exclude_unset=True, mode="json")

    incoming_email = str(updates.get("email") or "").strip().lower()
    if incoming_email:
        existing_email_user = db.query(PdaUser).filter(PdaUser.email == incoming_email, PdaUser.id != user_id).first()
        if existing_email_user:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email already in use")
        updates["email"] = incoming_email

    incoming_profile_name = updates.get("profile_name")
    if incoming_profile_name:
        existing_profile_user = db.query(PdaUser).filter(
            PdaUser.profile_name == incoming_profile_name,
            PdaUser.id != user_id
        ).first()
        if existing_profile_user:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name already in use")

    for field in ("name", "profile_name", "email", "phno", "dept", "dob", "gender", "is_member", "instagram_url", "linkedin_url", "github_url"):
        if field in updates:
            setattr(user, field, updates[field])
    if "photo_url" in updates:
        user.image_url = updates.get("photo_url")

    member = db.query(PdaTeam).filter(PdaTeam.user_id == user_id).order_by(PdaTeam.id.asc()).first()
    if bool(updates.get("clear_team")):
        db.query(PdaTeam).filter(PdaTeam.user_id == user_id).delete(synchronize_session=False)
        member = None
    elif ("team" in updates) or ("designation" in updates):
        next_team = updates.get("team")
        next_designation = updates.get("designation")
        if not member:
            if next_team is not None or next_designation is not None:
                member = PdaTeam(
                    user_id=user_id,
                    team=next_team,
                    designation=next_designation,
                )
                db.add(member)
                db.flush()
        else:
            if "team" in updates:
                member.team = next_team
            if "designation" in updates:
                member.designation = next_designation
            if member.team is None and member.designation is None:
                db.delete(member)
                member = None

    _sync_member_status_from_team(user, member)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Duplicate value violates a unique constraint")

    db.refresh(user)
    if member:
        db.refresh(member)
    else:
        member = db.query(PdaTeam).filter(PdaTeam.user_id == user_id).order_by(PdaTeam.id.asc()).first()

    log_admin_action(
        db,
        admin,
        "Update PDA user",
        request.method if request else None,
        request.url.path if request else None,
        {"user_id": user_id, "target_regno": user.regno},
    )
    recruit_state = get_recruitment_state(db, user.id, user=user)
    return _build_admin_user_response(user, member, recruit_state)


@router.post("/pda-admin/team", response_model=PdaTeamResponse)
def create_team_member(
    member_data: PdaTeamCreate,
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    payload = member_data.model_dump(mode="json")
    desired_profile_name = str(payload.get("profile_name") or "").strip().lower() or None
    if desired_profile_name and not is_profile_name_valid(desired_profile_name):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid profile name format")
    password_value = str(payload.get("password") or "").strip()
    user = None
    if payload.get("user_id"):
        user = db.query(PdaUser).filter(PdaUser.id == payload["user_id"]).first()
        if user:
            for field in ("name", "email", "phno", "dept", "instagram_url", "linkedin_url", "github_url"):
                if payload.get(field):
                    setattr(user, field, payload[field])
            if payload.get("dob"):
                user.dob = payload["dob"]
            if payload.get("gender"):
                user.gender = payload["gender"]
            if desired_profile_name and desired_profile_name != user.profile_name:
                existing_profile = db.query(PdaUser).filter(PdaUser.profile_name == desired_profile_name).first()
                if existing_profile and existing_profile.id != user.id:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name already in use")
                community_conflict = db.query(PersohubCommunity).filter(
                    PersohubCommunity.profile_id == desired_profile_name
                ).first()
                if community_conflict:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name reserved by community")
                user.profile_name = desired_profile_name
    if not user and payload.get("regno"):
        user = db.query(PdaUser).filter(PdaUser.regno == payload["regno"]).first()
        if not user:
            if desired_profile_name:
                existing_profile = db.query(PdaUser).filter(PdaUser.profile_name == desired_profile_name).first()
                if existing_profile:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name already in use")
                community_conflict = db.query(PersohubCommunity).filter(
                    PersohubCommunity.profile_id == desired_profile_name
                ).first()
                if community_conflict:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name reserved by community")
            user = PdaUser(
                regno=payload["regno"],
                email=payload.get("email") or f"{payload['regno']}@pda.com",
                hashed_password=get_password_hash(password_value or "password"),
                name=payload.get("name") or f"PDA Member {payload['regno']}",
                profile_name=desired_profile_name or generate_unique_profile_name(db, payload.get("name") or payload["regno"]),
                dob=payload.get("dob"),
                gender=payload.get("gender"),
                phno=payload.get("phno"),
                dept=payload.get("dept"),
                instagram_url=payload.get("instagram_url"),
                linkedin_url=payload.get("linkedin_url"),
                github_url=payload.get("github_url"),
                image_url=payload.get("photo_url"),
                json_content={},
                is_member=True
            )
            db.add(user)
            db.flush()
        else:
            for field in ("name", "email", "phno", "dept", "instagram_url", "linkedin_url", "github_url"):
                if payload.get(field):
                    setattr(user, field, payload[field])
            if payload.get("dob"):
                user.dob = payload["dob"]
            if payload.get("gender"):
                user.gender = payload["gender"]
            if desired_profile_name and desired_profile_name != user.profile_name:
                existing_profile = db.query(PdaUser).filter(PdaUser.profile_name == desired_profile_name).first()
                if existing_profile and existing_profile.id != user.id:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name already in use")
                community_conflict = db.query(PersohubCommunity).filter(
                    PersohubCommunity.profile_id == desired_profile_name
                ).first()
                if community_conflict:
                    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Profile name reserved by community")
                user.profile_name = desired_profile_name

    if not user:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="user_id or regno is required")

    new_member = PdaTeam(
        user_id=user.id,
        team=payload.get("team"),
        designation=payload.get("designation")
    )
    _sync_member_status_from_team(user, new_member)
    if payload.get("photo_url") and (not user.image_url or user.image_url != payload["photo_url"]):
        user.image_url = payload["photo_url"]
    db.add(new_member)
    db.commit()
    db.refresh(new_member)
    log_admin_action(
        db,
        admin,
        "Create team member",
        request.method if request else None,
        request.url.path if request else None,
        {"member_id": new_member.id, "user_id": user.id, "target_regno": user.regno},
    )
    resume_state = get_recruitment_state(db, user.id, user=user)
    return _build_team_response(new_member, user, resume_state.get("resume_url"))


@router.put("/pda-admin/team/{member_id}", response_model=PdaTeamResponse)
def update_team_member(
    member_id: int,
    member_data: PdaTeamUpdate,
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    member = db.query(PdaTeam).filter(PdaTeam.id == member_id).first()
    if not member:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Team member not found")

    updates = member_data.model_dump(exclude_unset=True, mode="json")
    for field in ("team", "designation"):
        if field in updates:
            setattr(member, field, updates[field])

    user = None
    if updates.get("user_id"):
        user = db.query(PdaUser).filter(PdaUser.id == updates["user_id"]).first()
    if not user and member.user_id:
        user = db.query(PdaUser).filter(PdaUser.id == member.user_id).first()
    if not user and updates.get("regno"):
        user = db.query(PdaUser).filter(PdaUser.regno == updates["regno"]).first()
    if user:
        member.user_id = user.id
        for field in ("name", "email", "phno", "dept", "dob", "instagram_url", "linkedin_url", "github_url"):
            if field in updates and updates[field] is not None:
                setattr(user, field, updates[field])
        if "photo_url" in updates and updates.get("photo_url"):
            if user.image_url != updates["photo_url"]:
                user.image_url = updates["photo_url"]

    _sync_member_status_from_team(user, member)

    db.commit()
    db.refresh(member)
    if not user and member.user_id:
        user = db.query(PdaUser).filter(PdaUser.id == member.user_id).first()
    meta = {"member_id": member_id}
    if user:
        meta.update({"user_id": user.id, "target_regno": user.regno})
    log_admin_action(
        db,
        admin,
        "Update team member",
        request.method if request else None,
        request.url.path if request else None,
        meta,
    )
    resume_url = None
    if user and user.id:
        resume_state = get_recruitment_state(db, user.id, user=user)
        resume_url = resume_state.get("resume_url")
    return _build_team_response(member, user, resume_url)


@router.delete("/pda-admin/team/{member_id}")
def delete_team_member(
    member_id: int,
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    member = db.query(PdaTeam).filter(PdaTeam.id == member_id).first()
    if not member:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Team member not found")
    user_id = member.user_id
    user = db.query(PdaUser).filter(PdaUser.id == user_id).first() if user_id else None
    db.delete(member)
    if user_id:
        remaining = db.query(PdaTeam).filter(PdaTeam.user_id == user_id).first()
        if not remaining:
            user = db.query(PdaUser).filter(PdaUser.id == user_id).first()
            if user:
                user.is_member = False
    db.commit()
    meta = {"member_id": member_id}
    if user_id:
        meta["user_id"] = user_id
    if user:
        meta["target_regno"] = user.regno
    log_admin_action(
        db,
        admin,
        "Delete team member",
        request.method if request else None,
        request.url.path if request else None,
        meta,
    )
    return {"message": "Team member deleted successfully"}


@router.delete("/pda-admin/users/{user_id}")
def delete_pda_user(
    user_id: int,
    force: bool = False,
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    user = db.query(PdaUser).filter(PdaUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    blocking = _user_dependency_checks(db, user_id)
    if blocking and not force:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"User has related records: {', '.join(blocking)}"
        )

    if force:
        team_ids_led = [t.id for t in db.query(PdaEventTeam).filter(PdaEventTeam.team_lead_user_id == user_id).all()]
        team_ids_member = [t.id for t in db.query(PdaEventTeamMember).filter(PdaEventTeamMember.user_id == user_id).all()]
        team_ids = list({*team_ids_led, *team_ids_member})

        if team_ids:
            db.query(PdaEventRegistration).filter(PdaEventRegistration.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventAttendance).filter(PdaEventAttendance.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventScore).filter(PdaEventScore.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventBadge).filter(PdaEventBadge.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventInvite).filter(PdaEventInvite.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventTeamMember).filter(PdaEventTeamMember.team_id.in_(team_ids)).delete(synchronize_session=False)
            db.query(PdaEventTeam).filter(PdaEventTeam.id.in_(team_ids_led)).delete(synchronize_session=False)

        db.query(PdaEventRegistration).filter(PdaEventRegistration.user_id == user_id).delete(synchronize_session=False)
        db.query(PdaEventTeamMember).filter(PdaEventTeamMember.user_id == user_id).delete(synchronize_session=False)
        db.query(PdaEventAttendance).filter(
            (PdaEventAttendance.user_id == user_id) | (PdaEventAttendance.marked_by_user_id == user_id)
        ).delete(synchronize_session=False)
        db.query(PdaEventScore).filter(PdaEventScore.user_id == user_id).delete(synchronize_session=False)
        db.query(PdaEventBadge).filter(PdaEventBadge.user_id == user_id).delete(synchronize_session=False)
        db.query(PdaEventInvite).filter(
            (PdaEventInvite.invited_user_id == user_id) | (PdaEventInvite.invited_by_user_id == user_id)
        ).delete(synchronize_session=False)

        post_ids = [p.id for p in db.query(PersohubPost).filter(PersohubPost.admin_id == user_id).all()]
        if post_ids:
            db.query(PersohubPostAttachment).filter(PersohubPostAttachment.post_id.in_(post_ids)).delete(synchronize_session=False)
            db.query(PersohubPostLike).filter(PersohubPostLike.post_id.in_(post_ids)).delete(synchronize_session=False)
            db.query(PersohubPostComment).filter(PersohubPostComment.post_id.in_(post_ids)).delete(synchronize_session=False)
            db.query(PersohubPostMention).filter(PersohubPostMention.post_id.in_(post_ids)).delete(synchronize_session=False)
            db.query(PersohubPostHashtag).filter(PersohubPostHashtag.post_id.in_(post_ids)).delete(synchronize_session=False)
            db.query(PersohubPost).filter(PersohubPost.id.in_(post_ids)).delete(synchronize_session=False)

        db.query(PersohubPostLike).filter(PersohubPostLike.user_id == user_id).delete(synchronize_session=False)
        db.query(PersohubPostComment).filter(PersohubPostComment.user_id == user_id).delete(synchronize_session=False)
        db.query(PersohubPostMention).filter(PersohubPostMention.user_id == user_id).delete(synchronize_session=False)
        db.query(PersohubCommunityFollow).filter(PersohubCommunityFollow.user_id == user_id).delete(synchronize_session=False)

        community_ids = [c.id for c in db.query(PersohubCommunity).filter(PersohubCommunity.admin_id == user_id).all()]
        if community_ids:
            community_posts = [p.id for p in db.query(PersohubPost).filter(PersohubPost.community_id.in_(community_ids)).all()]
            if community_posts:
                db.query(PersohubPostAttachment).filter(PersohubPostAttachment.post_id.in_(community_posts)).delete(synchronize_session=False)
                db.query(PersohubPostLike).filter(PersohubPostLike.post_id.in_(community_posts)).delete(synchronize_session=False)
                db.query(PersohubPostComment).filter(PersohubPostComment.post_id.in_(community_posts)).delete(synchronize_session=False)
                db.query(PersohubPostMention).filter(PersohubPostMention.post_id.in_(community_posts)).delete(synchronize_session=False)
                db.query(PersohubPostHashtag).filter(PersohubPostHashtag.post_id.in_(community_posts)).delete(synchronize_session=False)
                db.query(PersohubPost).filter(PersohubPost.id.in_(community_posts)).delete(synchronize_session=False)
            db.query(PersohubCommunityFollow).filter(PersohubCommunityFollow.community_id.in_(community_ids)).delete(synchronize_session=False)
            db.query(PersohubCommunity).filter(PersohubCommunity.id.in_(community_ids)).delete(synchronize_session=False)

    db.query(PdaTeam).filter(PdaTeam.user_id == user_id).delete(synchronize_session=False)
    db.query(PdaAdmin).filter(PdaAdmin.user_id == user_id).delete(synchronize_session=False)
    db.delete(user)
    db.commit()
    log_admin_action(
        db,
        admin,
        "Delete PDA user",
        request.method if request else None,
        request.url.path if request else None,
        {"user_id": user_id, "target_regno": user.regno},
    )
    return {"message": "User deleted successfully"}


@router.get("/pda-admin/team/export")
def export_team_members(
    format: str = "csv",
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    rows = (
        db.query(PdaTeam, PdaUser)
        .join(PdaUser, PdaTeam.user_id == PdaUser.id, isouter=True)
        .order_by(PdaTeam.team.asc().nullslast(), PdaTeam.designation.asc().nullslast(), PdaUser.name.asc().nullslast())
        .all()
    )
    if format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Regno", "Team", "Designation", "Email", "Phone"])
        for m, u in rows:
            ws.append([u.name if u else None, u.regno if u else None, m.team, m.designation, u.email if u else None, u.phno if u else None])
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        headers = {"Content-Disposition": "attachment; filename=team.xlsx"}
        return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Name", "Regno", "Team", "Designation", "Email", "Phone"])
    for m, u in rows:
        writer.writerow([u.name if u else None, u.regno if u else None, m.team, m.designation, u.email if u else None, u.phno if u else None])
    headers = {"Content-Disposition": "attachment; filename=team.csv"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers=headers)


@router.get("/pda-admin/users/export")
def export_pda_users(
    format: str = "csv",
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    users = (
        db.query(PdaUser)
        .order_by(PdaUser.name.asc().nullslast(), PdaUser.regno.asc())
        .all()
    )
    user_ids = [user.id for user in users]
    team_rows = (
        db.query(PdaTeam)
        .filter(PdaTeam.user_id.in_(user_ids))
        .order_by(PdaTeam.id.asc())
        .all()
        if user_ids else []
    )
    team_map = {}
    for member in team_rows:
        if member.user_id not in team_map:
            team_map[member.user_id] = member
    recruit_map = get_recruitment_state_map(db, users)

    def _row(user: PdaUser):
        recruit = recruit_map.get(user.id, {})
        member = team_map.get(user.id)
        batch = str(user.regno or "")[:4] if user.regno else None
        return [
            user.name,
            user.profile_name,
            user.regno,
            user.email,
            user.phno,
            user.dept,
            user.gender,
            batch,
            "YES" if user.is_member else "NO",
            "YES" if recruit.get("is_applied") else "NO",
            recruit.get("preferred_team"),
            member.team if member else None,
            member.designation if member else None,
            user.created_at.isoformat() if user.created_at else None,
        ]

    headers_row = [
        "Name",
        "Profile Name",
        "Regno",
        "Email",
        "Phone",
        "Department",
        "Gender",
        "Batch",
        "Is Member",
        "Is Applied",
        "Preferred Team",
        "Team",
        "Designation",
        "Created At",
    ]

    if format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.append(headers_row)
        for user in users:
            ws.append(_row(user))
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        headers = {"Content-Disposition": "attachment; filename=users.xlsx"}
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers,
        )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers_row)
    for user in users:
        writer.writerow(_row(user))
    headers = {"Content-Disposition": "attachment; filename=users.csv"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers=headers)


@router.post("/pda-admin/gallery", response_model=PdaGalleryResponse)
def create_gallery_item(
    gallery_data: PdaGalleryCreate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    new_item = PdaGallery(**gallery_data.model_dump())
    db.add(new_item)
    db.commit()
    db.refresh(new_item)
    log_admin_action(db, admin, "Create gallery item", request.method if request else None, request.url.path if request else None, {"gallery_id": new_item.id})
    return PdaGalleryResponse.model_validate(new_item)


@router.put("/pda-admin/gallery/{item_id}", response_model=PdaGalleryResponse)
def update_gallery_item(
    item_id: int,
    gallery_data: PdaGalleryUpdate,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    item = db.query(PdaGallery).filter(PdaGallery.id == item_id).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Gallery item not found")

    for field, value in gallery_data.model_dump(exclude_unset=True).items():
        setattr(item, field, value)

    db.commit()
    db.refresh(item)
    log_admin_action(db, admin, "Update gallery item", request.method if request else None, request.url.path if request else None, {"gallery_id": item_id})
    return PdaGalleryResponse.model_validate(item)


@router.delete("/pda-admin/gallery/{item_id}")
def delete_gallery_item(
    item_id: int,
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    item = db.query(PdaGallery).filter(PdaGallery.id == item_id).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Gallery item not found")
    db.delete(item)
    db.commit()
    log_admin_action(db, admin, "Delete gallery item", request.method if request else None, request.url.path if request else None, {"gallery_id": item_id})
    return {"message": "Gallery item deleted successfully"}


@router.post("/pda-admin/posters")
def upload_pda_poster(
    file: UploadFile = File(...),
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing filename")
    allowed_types = ["image/png", "image/jpeg", "image/webp"]
    url = _upload_to_s3(file, "posters", allowed_types=allowed_types)
    log_admin_action(db, admin, "Upload PDA poster", request.method if request else None, request.url.path if request else None, {"file": file.filename})
    return {"url": url}


@router.post("/pda-admin/posters/presign", response_model=PresignResponse)
def presign_pda_poster(
    payload: PresignRequest,
    admin: PdaUser = Depends(require_pda_home_admin)
):
    return _generate_presigned_put_url(
        "posters",
        payload.filename,
        payload.content_type,
        allowed_types=["image/png", "image/jpeg", "image/webp", "application/pdf"]
    )


@router.post("/pda-admin/posters/pdf-preview", response_model=PdaPdfPreviewGenerateResponse)
def generate_pda_pdf_preview_images(
    payload: PdaPdfPreviewGenerateRequest,
    _: PdaUser = Depends(require_pda_home_admin),
):
    if not S3_CLIENT or not S3_BUCKET_NAME:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="S3 not configured")

    source_key = _extract_s3_key_from_url(payload.s3_url)
    if not source_key:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid S3 URL")
    if not source_key.startswith("posters/"):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="File does not belong to posters bucket path")
    if not source_key.lower().endswith(".pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only PDF files are supported")

    try:
        source_obj = S3_CLIENT.get_object(Bucket=S3_BUCKET_NAME, Key=source_key)
        pdf_bytes = source_obj["Body"].read()
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Failed to read source PDF") from exc

    preview_bytes = _render_pdf_preview_images(pdf_bytes, max_pages=payload.max_pages)
    if not preview_bytes:
        return PdaPdfPreviewGenerateResponse(preview_image_urls=[], pages_generated=0)

    source_stem = Path(source_key).stem
    run_token = os.urandom(6).hex()
    preview_urls: List[str] = []
    for idx, image_bytes in enumerate(preview_bytes, start=1):
        preview_key = f"posters/previews/{source_stem}-{run_token}/page_{idx:03d}.webp"
        try:
            S3_CLIENT.put_object(
                Bucket=S3_BUCKET_NAME,
                Key=preview_key,
                Body=image_bytes,
                ContentType="image/webp",
                CacheControl="public, max-age=31536000, immutable",
            )
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Failed to upload preview image") from exc
        preview_urls.append(_build_s3_url(preview_key))

    return PdaPdfPreviewGenerateResponse(
        preview_image_urls=preview_urls,
        pages_generated=len(preview_urls),
    )


@router.post("/pda-admin/gallery-uploads")
def upload_pda_gallery_image(
    file: UploadFile = File(...),
    admin: PdaUser = Depends(require_pda_home_admin),
    db: Session = Depends(get_db),
    request: Request = None
):
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing filename")
    allowed_types = ["image/png", "image/jpeg", "image/webp"]
    url = _upload_to_s3(file, "gallery", allowed_types=allowed_types)
    log_admin_action(db, admin, "Upload gallery image", request.method if request else None, request.url.path if request else None, {"file": file.filename})
    return {"url": url}


@router.post("/pda-admin/gallery-uploads/presign", response_model=PresignResponse)
def presign_pda_gallery_image(
    payload: PresignRequest,
    admin: PdaUser = Depends(require_pda_home_admin)
):
    return _generate_presigned_put_url(
        "gallery",
        payload.filename,
        payload.content_type,
        allowed_types=["image/png", "image/jpeg", "image/webp"]
    )


@router.post("/pda-admin/team-uploads")
def upload_pda_team_image(
    file: UploadFile = File(...),
    admin: PdaUser = Depends(require_superadmin),
    db: Session = Depends(get_db),
    request: Request = None
):
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing filename")
    allowed_types = ["image/png", "image/jpeg", "image/webp"]
    url = _upload_to_s3(file, "team", allowed_types=allowed_types)
    log_admin_action(db, admin, "Upload team image", request.method if request else None, request.url.path if request else None, {"file": file.filename})
    return {"url": url}


@router.post("/pda-admin/team-uploads/presign", response_model=PresignResponse)
def presign_pda_team_image(
    payload: PresignRequest,
    admin: PdaUser = Depends(require_superadmin)
):
    return _generate_presigned_put_url(
        "team",
        payload.filename,
        payload.content_type,
        allowed_types=["image/png", "image/jpeg", "image/webp"]
    )
